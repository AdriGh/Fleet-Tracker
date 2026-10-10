# -*- coding: utf-8 -*-
"""Tests de Reports v2.22: guardados, favoritos, envíos programados,
adjuntos del mailer, asistente de IA y rutas/RBAC.

`python backend/tests/test_reports_saved.py`. DB SQLite temporal fresca.
SIN red: el mailer (SMTP) y la API de Anthropic se reemplazan por fakes.

Lo que NO hay que perder:
  - guardados: 'mine' solo lo ve/edita su dueño; 'org' lo ve toda la org y
    crearlo/editarlo ajeno exige maint.edit (403); spec validada por el
    motor (400 con mensaje); borrar quita sus ★; resolve(ref) respeta la
    visibilidad salvo en el scheduler;
  - favoritos idempotentes, refs inválidos -> 400;
  - schedules: next_run (daily/weekly/monthly, fin de mes, DST de la org),
    validaciones, RECLAMO atómico (un solo worker envía), estados sent /
    dry_run / error, adjunto CSV/XLSX, PII SIEMPRE enmascarada, recorre
    todas las orgs con su tenant;
  - mailer: adjuntos en multipart/mixed;
  - asistente: 503 sin clave / apagado; structured output -> spec validada;
    reintento con el error del validador; saneo; rechazo -> sin spec y con
    aviso; NUNCA manda filas de datos; fallbacks + caché + effort en la
    request real (cliente falso); rate limit;
  - RBAC (_scope_for) y orden de rutas.
"""
import asyncio
import email
import os
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

_BACKEND = Path(__file__).resolve().parent.parent
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

_TMP = Path(tempfile.mkdtemp())
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP / 'reports_saved_test.db'}"
os.environ.pop("ANTHROPIC_API_KEY", None)

from fastapi import HTTPException                               # noqa: E402
from sqlalchemy import update                                   # noqa: E402

from app import db                                              # noqa: E402
from app.api import routes                                      # noqa: E402
from app.core import (                                          # noqa: E402
    customers, mailer, org_config, ratelimit, report_assistant as ra,
    report_engine as eng, report_saved as saved, report_schedules as sch,
    tenant, workorders,
)
from app.main import _scope_for                                 # noqa: E402

ORG = db.default_org_id()
tenant.set_current_org(ORG)

# El mailer NUNCA sale a la red: settings "no configurado" (= simulado) y,
# cuando un test quiere "enviar", se reemplaza send_email por un fake.
mailer.load_settings = lambda: mailer.EmailSettings({})
_REAL_SEND = mailer.send_email


def _req(role: str, uid: int, org: int = ORG):
    return SimpleNamespace(
        state=SimpleNamespace(user={"id": uid, "org_id": org, "role": role,
                                    "name": f"Test {role}",
                                    "username": f"{role}{uid}"}),
        client=SimpleNamespace(host="127.0.0.1"), headers={})


MECH = _req("mechanic", 1)          # maint.edit, sin pii.view
VIEW = _req("viewer", 2)            # sin scopes
DISP = _req("dispatcher", 3)        # maint.edit + pii.view


def _http(fn, status: int, *args, **kw):
    try:
        res = fn(*args, **kw)
        if asyncio.iscoroutine(res):
            asyncio.run(res)
    except HTTPException as exc:
        assert exc.status_code == status, (exc.status_code, exc.detail)
        return exc
    raise AssertionError(f"debió dar {status}")


def _raises(fn, needle: str, *args, exc=ValueError, **kw):
    try:
        res = fn(*args, **kw)
        if asyncio.iscoroutine(res):
            asyncio.run(res)
    except exc as e:
        assert needle in str(e), f"{needle!r} not in {e!s}"
        return
    raise AssertionError(f"debió lanzar {exc.__name__} con {needle!r}")


SPEC = {"source": "work_orders",
        "columns": [{"field": "display_no"}, {"field": "unit"},
                    {"field": "total"}]}

_seeded = {}


def _seed():
    if _seeded:
        return
    w = workorders.create_wo("MEM 77", "ZZSECRET brake job",
                             mechanic="Luis")
    workorders.add_line(w["id"], "part", "brake pad", 2, 40)
    c = customers.create_customer({"name": "Juan Perez",
                                   "email": "juan@example.com",
                                   "phone": "901-555-7777"})
    _seeded.update(wo=w["id"], customer=c["id"])


# ----- Guardados + favoritos ---------------------------------------------------------

def test_saved_reports_visibility_and_permissions():
    _seed()
    mine = routes.reports_saved_create(routes.SavedReportIn(
        name="  My   WOs ", spec=SPEC, based_on="sys:work-order-open-summary"),
        VIEW)
    assert mine["name"] == "My WOs" and mine["visibility"] == "mine"
    assert mine["based_on"] == "sys:work-order-open-summary" \
        and mine["mine"]
    # Un viewer no comparte con la org.
    _http(routes.reports_saved_create, 403, routes.SavedReportIn(
        name="Shared", spec=SPEC, visibility="org"), VIEW)
    org_r = routes.reports_saved_create(routes.SavedReportIn(
        name="Shared", spec=SPEC, visibility="org"), MECH)
    # Spec inválida / based_on inexistente / visibilidad rara -> 400.
    e = _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="Bad", spec={"source": "work_orders",
                          "columns": [{"field": "nope"}]}), VIEW)
    assert "unknown field 'nope'" in e.detail
    _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="Bad", spec=SPEC, based_on="sys:nope"), VIEW)
    _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="Bad", spec=SPEC, visibility="public"), MECH)
    _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="   ", spec=SPEC), MECH)
    # Visibilidad.
    lib_v = routes.reports_library(VIEW)
    names_v = {s["name"] for s in lib_v["saved"]}
    assert names_v == {"My WOs", "Shared"}
    assert len(lib_v["system"]) >= 30
    names_m = {s["name"] for s in routes.reports_library(MECH)["saved"]}
    assert names_m == {"Shared"}
    _http(routes.reports_saved_get, 404, mine["id"], MECH)
    assert routes.reports_saved_get(mine["id"], VIEW)["spec"]["source"] == \
        "work_orders"
    # Editar: el dueño sí; un 'org' ajeno exige maint.edit.
    upd = routes.reports_saved_patch(mine["id"], routes.SavedReportPatch(
        description="desc"), VIEW)
    assert upd["description"] == "desc"
    _http(routes.reports_saved_patch, 403, org_r["id"],
          routes.SavedReportPatch(name="Hacked"), VIEW)
    _http(routes.reports_saved_patch, 403, mine["id"],
          routes.SavedReportPatch(visibility="org"), VIEW)
    ok = routes.reports_saved_patch(org_r["id"], routes.SavedReportPatch(
        name="Shared v2", spec={**SPEC, "totals": True}), DISP)
    assert ok["name"] == "Shared v2" and ok["can_edit"]
    _http(routes.reports_saved_patch, 403, org_r["id"],
          routes.SavedReportPatch(visibility="mine"), DISP)
    _http(routes.reports_saved_patch, 404, mine["id"],
          routes.SavedReportPatch(name="x"), MECH)
    # resolve: visibilidad salvo en modo sistema.
    _raises(saved.resolve, "unknown report", f"saved:{mine['id']}",
            user_id=1)
    spec, meta = saved.resolve(f"saved:{mine['id']}", system=True)
    assert spec["source"] == "work_orders" and meta["kind"] == "saved"
    _raises(saved.resolve, "can't override 'source'",
            f"saved:{mine['id']}", overrides={"source": "parts"}, user_id=2)
    # Run por ref con overrides.
    res = asyncio.run(routes.reports_run(routes.ReportRunIn(
        ref=f"saved:{mine['id']}",
        overrides={"time": {"field": "created_at", "preset": "today"}}),
        VIEW))
    assert res["report"]["name"] == "My WOs" and res["row_count"] == 1
    _http(routes.reports_run, 400, routes.ReportRunIn(
        ref=f"saved:{mine['id']}"), MECH)
    _http(routes.reports_run, 400, routes.ReportRunIn(), MECH)
    _http(routes.reports_run, 400, routes.ReportRunIn(
        spec=SPEC, overrides={"limit": 1}), MECH)
    res = asyncio.run(routes.reports_run(routes.ReportRunIn(
        ref="sys:work-order-open-summary", version="waiting_parts"), VIEW))
    assert res["report"]["version"] == "waiting_parts"
    # Favoritos.
    fav = routes.reports_favorite_add(routes.ReportFavoriteIn(
        ref="sys:cost-per-mile"), VIEW)
    fav = routes.reports_favorite_add(routes.ReportFavoriteIn(
        ref="sys:cost-per-mile"), VIEW)
    assert fav["favorites"] == ["sys:cost-per-mile"]
    fav = routes.reports_favorite_add(routes.ReportFavoriteIn(
        ref=f"saved:{mine['id']}"), VIEW)
    assert len(fav["favorites"]) == 2
    _http(routes.reports_favorite_add, 400, routes.ReportFavoriteIn(
        ref="sys:nope"), VIEW)
    _http(routes.reports_favorite_add, 400, routes.ReportFavoriteIn(
        ref=f"saved:{mine['id']}"), MECH)          # no lo ve
    _http(routes.reports_favorite_add, 400, routes.ReportFavoriteIn(
        ref="cost-per-mile"), VIEW)
    assert routes.reports_library(MECH)["favorites"] == []
    out = routes.reports_favorite_remove(VIEW, ref="sys:cost-per-mile")
    assert out["favorites"] == [f"saved:{mine['id']}"]
    _http(routes.reports_favorite_remove, 400, VIEW)
    # Borrar quita sus favoritos.
    routes.reports_saved_delete(mine["id"], VIEW)
    assert routes.reports_library(VIEW)["favorites"] == []
    _http(routes.reports_saved_delete, 404, mine["id"], VIEW)
    _http(routes.reports_saved_delete, 403, org_r["id"], VIEW)
    print("OK guardados (mine/org, permisos, resolve) + favoritos")


def test_run_route_pii_and_export():
    _seed()
    spec = {"source": "customers",
            "columns": [{"field": "name"}, {"field": "email"},
                        {"field": "phone"}]}
    r = asyncio.run(routes.reports_run(routes.ReportRunIn(spec=spec), MECH))
    assert "•••" in r["rows"][0]["email"] and r["pii_masked"]
    r = asyncio.run(routes.reports_run(routes.ReportRunIn(spec=spec), DISP))
    assert r["rows"][0]["email"] == "juan@example.com"
    _http(routes.reports_run, 400, routes.ReportRunIn(spec={
        **spec, "filters": {"items": [{"field": "email", "op": "contains",
                                       "value": "@"}]}}), MECH)
    resp = asyncio.run(routes.reports_export(routes.ReportExportIn(
        ref="sys:customer-balances", format="xlsx"), MECH))
    assert "attachment" in resp.headers["content-disposition"]
    assert "customer-balances" in resp.headers["content-disposition"]
    assert resp.body[:2] == b"PK"
    resp = asyncio.run(routes.reports_export(routes.ReportExportIn(
        spec=spec, format="csv"), MECH))
    assert resp.body.decode("utf-8").count("•••") >= 2
    _http(routes.reports_export, 400, routes.ReportExportIn(
        spec=spec, format="pdf"), MECH)
    src = routes.reports_sources()
    assert {s["key"] for s in src["sources"]} >= {"work_orders", "cpm"}
    assert src["limits"]["max_rows"] == 50000
    sysr = routes.reports_system("cost-per-mile")
    assert sysr["spec"]["source"] == "cpm" and sysr["dashboard"] == "cpm"
    _http(routes.reports_system, 404, "nope")
    print("OK run (PII por permiso) + export CSV/XLSX + sources/system")


# ----- Schedules -------------------------------------------------------------------------

def test_next_run_math():
    tz = ZoneInfo("America/Chicago")
    at = datetime(2026, 10, 9, 6, 30, tzinfo=tz)          # viernes
    assert sch.next_run("daily", 7, None, None, at, tz) == \
        datetime(2026, 10, 9, 7, 0, tzinfo=tz)
    assert sch.next_run("daily", 6, None, None, at, tz) == \
        datetime(2026, 10, 10, 6, 0, tzinfo=tz)
    assert sch.next_run("weekly", 7, 0, None, at, tz) == \
        datetime(2026, 10, 12, 7, 0, tzinfo=tz)            # lunes
    assert sch.next_run("weekly", 7, 4, None, at, tz) == \
        datetime(2026, 10, 9, 7, 0, tzinfo=tz)             # hoy viernes
    assert sch.next_run("weekly", 6, 4, None, at, tz) == \
        datetime(2026, 10, 16, 6, 0, tzinfo=tz)
    assert sch.next_run("monthly", 7, None, 31,
                        datetime(2027, 2, 1, tzinfo=tz), tz) == \
        datetime(2027, 2, 28, 7, 0, tzinfo=tz)              # recorte
    assert sch.next_run("monthly", 7, None, 5, at, tz) == \
        datetime(2026, 11, 5, 7, 0, tzinfo=tz)
    assert sch.next_run("monthly", 7, None, 1,
                        datetime(2026, 12, 15, tzinfo=tz), tz) == \
        datetime(2027, 1, 1, 7, 0, tzinfo=tz)
    # DST (1-nov-2026 en Chicago): sigue siendo las 7:00 locales.
    nx = sch.next_run("daily", 7, None, None,
                      datetime(2026, 10, 31, 8, 0, tzinfo=tz), tz)
    assert nx.hour == 7 and nx.utcoffset() == timedelta(hours=-6)
    print("OK próximo envío (diario/semanal/mensual, fin de mes, DST)")


def _capture():
    sent = []

    def fake_send(settings, to, cc, subject, body, html="",
                  attachments=None):
        sent.append({"to": to, "cc": cc, "subject": subject, "body": body,
                     "attachments": attachments or []})
        return {"ok": True, "error": "", "simulated": False}
    return sent, fake_send


def test_schedules_crud_claim_and_send():
    _seed()
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="not-an-email"), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients=[]), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="a@b.co",
        cadence="hourly"), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="a@b.co", hour=24), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="a@b.co",
        cadence="weekly", weekday=7), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="a@b.co",
        cadence="monthly", day_of_month=0), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:nope", recipients="a@b.co"), MECH)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", version="nope",
        recipients="a@b.co"), MECH)
    s1 = routes.reports_schedules_create(routes.ReportScheduleIn(
        report_ref="sys:customer-balances", cadence="daily", hour=7,
        recipients="ops@example.com; boss@example.com, ops@example.com",
        format="csv"), MECH)
    assert s1["recipients"] == ["ops@example.com", "boss@example.com"]
    assert s1["name"].startswith("Customer Balances")
    assert s1["when"] == "Daily at 07:00" and s1["active"]
    assert s1["next_run_at"] > org_config.org_now().replace(
        tzinfo=None).isoformat()
    assert s1["created_by"] == "Test mechanic"
    # Pausar / reanudar / cambiar cadencia.
    p = routes.reports_schedules_patch(s1["id"], routes.ReportSchedulePatch(
        active=False), MECH)
    assert p["next_run_at"] is None
    p = routes.reports_schedules_patch(s1["id"], routes.ReportSchedulePatch(
        active=True, cadence="weekly", weekday=2), MECH)
    assert p["when"] == "Every Wednesday at 07:00" and p["next_run_at"]
    _http(routes.reports_schedules_patch, 400, s1["id"],
          routes.ReportSchedulePatch(format="pdf"), MECH)
    _http(routes.reports_schedules_patch, 404, 999,
          routes.ReportSchedulePatch(name="x"), MECH)
    assert [s["id"] for s in routes.reports_schedules_list()["schedules"]] \
        == [s1["id"]]

    # Vencido -> maybe_run lo RECLAMA y envía (adjunto CSV, PII oculta).
    with db.SessionLocal() as s:
        s.execute(update(db.ReportSchedule).where(
            db.ReportSchedule.id == s1["id"]).values(
            next_run_at=datetime.now().replace(microsecond=0)
            - timedelta(minutes=5)))
        s.commit()
    sent, fake = _capture()
    mailer.send_email = fake
    try:
        assert asyncio.run(sch.maybe_run()) == 1
        assert asyncio.run(sch.maybe_run()) == 0     # ya avanzado
    finally:
        mailer.send_email = _REAL_SEND
    assert len(sent) == 1
    m = sent[0]
    assert m["to"] == "ops@example.com" and m["cc"] == ["boss@example.com"]
    assert m["subject"].startswith("Rigsmith report: Customer Balances")
    fname, data, mime = m["attachments"][0]
    assert fname.endswith(".csv") and mime.startswith("text/csv")
    assert "Juan Perez" in data.decode("utf-8")
    assert "Contact details are hidden" in m["body"]
    row = routes.reports_schedules_list()["schedules"][0]
    assert row["last_status"] == "sent" and row["last_run_at"]
    assert row["next_run_at"] > org_config.org_now().replace(
        tzinfo=None).isoformat()

    # Reclamo atómico: dos workers con la MISMA lectura -> uno solo gana.
    with db.SessionLocal() as s:
        s.execute(update(db.ReportSchedule).where(
            db.ReportSchedule.id == s1["id"]).values(
            next_run_at=datetime(2026, 1, 1, 7, 0)))
        s.commit()
        stale = s.get(db.ReportSchedule, s1["id"])
        s.expunge(stale)
    tz = org_config.org_tz()
    assert sch._claim(ORG, stale, tz) is True
    assert sch._claim(ORG, stale, tz) is False

    # Sin mailer configurado -> dry_run (y se dice que no salió nada).
    out = asyncio.run(routes.reports_schedules_send_now(s1["id"]))
    assert out["last_status"] == "dry_run" and "nothing was sent" in \
        out["last_error"]
    # XLSX + PII de un reporte de customers con email.
    sp = routes.reports_saved_create(routes.SavedReportIn(
        name="Contacts", visibility="org",
        spec={"source": "customers",
              "columns": [{"field": "name"}, {"field": "email"}]}), DISP)
    s2 = routes.reports_schedules_create(routes.ReportScheduleIn(
        report_ref=f"saved:{sp['id']}", recipients="a@b.co",
        format="xlsx", cadence="monthly", day_of_month=31), DISP)
    assert s2["when"] == "Monthly on day 31 at 07:00"
    sent, fake = _capture()
    mailer.send_email = fake
    try:
        out = asyncio.run(sch.send_now(s2["id"]))
    finally:
        mailer.send_email = _REAL_SEND
    assert out["last_status"] == "sent" and out["rows"] == 1
    from io import BytesIO
    from openpyxl import load_workbook
    fname, data, mime = sent[0]["attachments"][0]
    assert fname.endswith(".xlsx")
    cells = [c.value for row in load_workbook(BytesIO(data)).active
             for c in row]
    assert "juan@example.com" not in cells
    assert any(isinstance(v, str) and "•••" in v for v in cells)
    # Borrar el reporte guardado se lleva su schedule (cascada).
    routes.reports_saved_delete(sp["id"], DISP)
    _http(routes.reports_schedules_send_now, 404, s2["id"])
    # Error: schedule huérfano (de antes de la cascada) -> 'error' claro.
    s3 = routes.reports_schedules_create(routes.ReportScheduleIn(
        report_ref="sys:cost-per-mile", recipients="a@b.co"), DISP)
    with db.SessionLocal() as ses:
        ses.execute(update(db.ReportSchedule).where(
            db.ReportSchedule.id == s3["id"]).values(
            report_ref=f"saved:{sp['id']}"))
        ses.commit()
    out = asyncio.run(routes.reports_schedules_send_now(s3["id"]))
    assert out["last_status"] == "error" and "unknown report" in \
        out["last_error"]
    # Un schedule de un 'mine' ajeno no se puede crear.
    mine = saved.create_saved({"name": "Private", "spec": SPEC}, user_id=9)
    _http(routes.reports_schedules_create, 400, routes.ReportScheduleIn(
        report_ref=f"saved:{mine['id']}", recipients="a@b.co"), MECH)
    routes.reports_schedules_delete(s3["id"])
    _http(routes.reports_schedules_delete, 404, s3["id"])
    print("OK schedules: validación, reclamo atómico, sent/dry_run/error, "
          "CSV/XLSX, PII oculta")


def test_scheduler_iterates_orgs():
    with db.SessionLocal() as s:
        o = db.Organization(slug="sched-org", name="Other", active=True,
                            created_at=datetime.now())
        s.add(o)
        s.commit()
        other = o.id
    tok = tenant.set_current_org(other)
    try:
        workorders.create_wo("OTHER 1", "other org wo")
        workorders.create_wo("OTHER 2", "other org wo")
        sp = saved.create_saved({"name": "All WOs", "visibility": "org",
                                 "spec": SPEC}, user_id=50, can_manage=True)
        s = sch.create_schedule({"report_ref": f"saved:{sp['id']}",
                                 "recipients": "o@x.co", "cadence": "daily"})
    finally:
        tenant.reset_current_org(tok)
    with db.SessionLocal() as ses:
        ses.execute(update(db.ReportSchedule).where(
            db.ReportSchedule.id == s["id"]).values(
            next_run_at=datetime(2026, 1, 1, 7, 0)))
        ses.commit()
    sent, fake = _capture()
    mailer.send_email = fake
    # El loop de alertas deja fijada la org 'default': maybe_run igual
    # procesa la otra org con SU tenant (2 WOs de esa org, no las nuestras).
    try:
        assert asyncio.run(sch.maybe_run()) == 1
    finally:
        mailer.send_email = _REAL_SEND
    assert "2 row(s)" in sent[0]["body"], sent[0]["body"]
    assert tenant.get_current_org() == ORG
    # Borrar el reporte guardado se lleva sus envíos programados (si no,
    # un reporte nuevo que reuse el id heredaría el schedule huérfano).
    tok = tenant.set_current_org(other)
    try:
        assert saved.delete_saved(sp["id"], 50, can_manage=True) is True
        with db.SessionLocal() as ses:
            assert ses.get(db.ReportSchedule, s["id"]) is None
    finally:
        tenant.reset_current_org(tok)
    print("OK scheduler: recorre todas las orgs con su tenant; borrar el "
          "reporte borra sus schedules")


def test_mailer_attachments():
    captured = {}

    class FakeSMTP:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def login(self, *a):
            captured["login"] = a

        def send_message(self, msg, to_addrs=None):
            captured["msg"] = msg.as_bytes()
            captured["to"] = to_addrs

    real = mailer.smtplib.SMTP_SSL
    mailer.smtplib.SMTP_SSL = FakeSMTP
    try:
        st = mailer.EmailSettings({"gmail_sender": "me@x.co",
                                   "gmail_app_password": "abcd efgh",
                                   "dry_run": False})
        res = mailer.send_email(st, "a@b.co", ["c@d.co"], "Report", "body",
                                attachments=[("r.csv", b"a,b\n1,2\n",
                                              "text/csv; charset=utf-8")])
    finally:
        mailer.smtplib.SMTP_SSL = real
    assert res == {"ok": True, "error": "", "simulated": False}
    msg = email.message_from_bytes(captured["msg"])
    assert msg.get_content_type() == "multipart/mixed"
    att = [p for p in msg.walk() if p.get_filename() == "r.csv"]
    assert att and att[0].get_content_type() == "text/csv"
    assert att[0].get_payload(decode=True) == b"a,b\n1,2\n"
    assert captured["to"] == ["a@b.co", "c@d.co"]
    print("OK mailer: adjuntos multipart/mixed")


# ----- Asistente de IA ---------------------------------------------------------------------

def _answer(spec: dict, explanation="Shows work orders.") -> ra.AssistantAnswer:
    base = {"filters": {"match": "all", "items": []}, "time": None,
            "group_by": [], "sort": [], "limit": None, "totals": False,
            "chart": None}
    full = {**base, **spec}
    full["columns"] = [{"agg": None, "label": None, **c}
                       for c in full["columns"]]
    return ra.AssistantAnswer.model_validate(
        {"spec": full, "explanation": explanation})


def test_assistant_unavailable_and_settings():
    real_key = ra._api_key
    ra._api_key = lambda: ("", "")
    try:
        st = routes.reports_assistant_status()
        assert st["available"] is False and st["message"] == \
            ra.NOT_CONFIGURED
        assert st["models"][0] == "claude-opus-5-5"
        e = _http(routes.reports_assistant, 503,
                  routes.ReportAssistantIn(prompt="x"), MECH)
        assert e.detail == "AI assistant isn't configured"
    finally:
        ra._api_key = real_key
    ra._api_key = lambda: ("sk-test", "env")
    try:
        org_config.save({"ai_assistant_enabled": False})
        e = _http(routes.reports_assistant, 503,
                  routes.ReportAssistantIn(prompt="x"), MECH)
        assert "turned off" in e.detail
        org_config.save({"ai_assistant_enabled": True,
                         "assistant_model": "claude-sonnet-5-5"})
        assert ra.status()["model"] == "claude-sonnet-5-5"
        assert ra.status()["available"] is True
        _raises(org_config.save, "invalid assistant model",
                {"assistant_model": "gpt-4o"})
        org_config.save({"assistant_model": ""})
        assert ra.status()["model"] == ra.DEFAULT_MODEL
    finally:
        ra._api_key = real_key
    print("OK asistente: 503 sin clave / apagado + settings de modelo")


def test_assistant_flow_with_fake_model():
    _seed()
    calls = []
    script = []

    async def fake_call(api_key, model, user_text):
        calls.append({"key": api_key, "model": model, "text": user_text})
        return script.pop(0)

    real_call, real_key = ra._call_model, ra._api_key
    ra._call_model = fake_call
    ra._api_key = lambda: ("sk-test", "env")
    try:
        good = _answer({"source": "wo_lines",
                        "columns": [{"field": "unit"},
                                    {"field": "amount", "agg": "sum"}],
                        "group_by": ["unit"],
                        "sort": [{"field": "sum_amount", "dir": "desc"}],
                        "time": {"field": "date", "preset": "last_90d",
                                 "date_from": None, "date_to": None},
                        "totals": True,
                        "chart": {"type": "bar", "x": "unit",
                                  "y": ["sum_amount"]}},
                       "Parts and labor cost by unit, last 90 days.")
        script.append(("end_turn", good, "claude-opus-5-5"))
        out = asyncio.run(routes.reports_assistant(routes.ReportAssistantIn(
            prompt="Parts cost by unit last 90 days"), MECH))
        assert out["ok"] and out["spec"]["source"] == "wo_lines"
        assert out["spec"]["time"]["preset"] == "last_90d"
        assert out["explanation"].startswith("Parts and labor")
        assert out["warnings"] == []
        eng.validate(out["spec"])
        t = calls[-1]["text"]
        assert "Request: Parts cost by unit last 90 days" in t
        assert org_config.org_today().isoformat() in t
        # Nunca viajan datos de la org (ni en el prompt ni en el system).
        assert "ZZSECRET" not in t and "ZZSECRET" not in ra.system_prompt()
        assert "Juan Perez" not in t and "Juan" not in ra.system_prompt()
        # Refinar: viaja la spec actual.
        script.append(("end_turn", good, "claude-opus-5-5"))
        asyncio.run(ra.ask("only DAL units", spec=out["spec"]))
        assert '"source": "wo_lines"' in calls[-1]["text"]
        # Inválida -> reintento con el error del validador.
        bad = _answer({"source": "wo_lines",
                       "columns": [{"field": "unit"},
                                   {"field": "amount", "agg": "sum"}],
                       "group_by": ["kind"]})
        script += [("end_turn", bad, "claude-opus-5-5"),
                   ("end_turn", good, "claude-opus-5-5")]
        out = asyncio.run(ra.ask("cost by unit"))
        assert out["ok"] and out["spec"]["group_by"] == ["unit"]
        assert "failed validation" in calls[-1]["text"]
        assert "must be in group_by" in calls[-1]["text"]
        # Inválida dos veces -> saneo (columna inventada fuera, aviso).
        weird = _answer({"source": "work_orders",
                         "columns": [{"field": "unit"},
                                     {"field": "made_up"}],
                         "sort": [{"field": "made_up", "dir": "asc"}]})
        script += [("end_turn", weird, "claude-opus-5-5"),
                   ("end_turn", weird, "claude-opus-5-5")]
        out = asyncio.run(ra.ask("units"))
        assert out["ok"] and out["spec"]["columns"] == [{"field": "unit"}]
        assert any("don't exist" in w for w in out["warnings"])
        # Rechazo / corte -> sin spec, con aviso claro.
        script.append(("refusal", None, "claude-opus-5-5"))
        out = asyncio.run(ra.ask("something"))
        assert out["spec"] is None and not out["ok"]
        assert "declined" in out["warnings"][0]
        script.append(("max_tokens", None, "claude-opus-5-5"))
        out = asyncio.run(ra.ask("something"))
        assert "cut off" in out["warnings"][0]
        # Respondió el modelo de fallback -> se avisa.
        script.append(("end_turn", good, "claude-opus-4-8"))
        out = asyncio.run(ra.ask("cost by unit"))
        assert out["model"] == "claude-opus-4-8"
        assert any("fallback" in w for w in out["warnings"])
        # PII: sin pii.view no puede filtrar por email -> saneo lo quita.
        pii = _answer({"source": "customers",
                       "columns": [{"field": "name"}],
                       "filters": {"match": "all", "items": [
                           {"field": "email", "op": "contains",
                            "value": "gmail"}]}})
        script += [("end_turn", pii, "m"), ("end_turn", pii, "m")]
        out = asyncio.run(ra.ask("customers with gmail", mask_pii=True))
        assert out["ok"] and out["spec"]["filters"]["items"] == []
        # Pedido vacío -> 400.
        _http(routes.reports_assistant, 400,
              routes.ReportAssistantIn(prompt="   "), MECH)
        # Rate limit por usuario (como el escaneo).
        ratelimit._reset_all()
        real_rate = ratelimit.SCAN_RATE
        ratelimit.SCAN_RATE = 1
        try:
            script.append(("end_turn", good, "m"))
            asyncio.run(routes.reports_assistant(routes.ReportAssistantIn(
                prompt="x"), VIEW))
            _http(routes.reports_assistant, 429,
                  routes.ReportAssistantIn(prompt="x"), VIEW)
        finally:
            ratelimit.SCAN_RATE = real_rate
            ratelimit._reset_all()
    finally:
        ra._call_model, ra._api_key = real_call, real_key
    print("OK asistente: spec validada, reintento, saneo, rechazos, sin "
          "datos al modelo, rate limit")


def test_assistant_request_shape():
    """La request real (cliente falso): structured output sin recursión,
    caché del system, thinking adaptive + effort, fallbacks default."""
    captured = []

    class FakeResp:
        def __init__(self, stop, parsed):
            self.stop_reason = stop
            self.parsed_output = parsed
            self.model = "claude-opus-5-5"

    class FakeMessages:
        def __init__(self, stop):
            self.stop = stop

        async def parse(self, **kw):
            captured.append(kw)
            if self.stop == "invalid":
                raise ValueError("bad json")
            return FakeResp(self.stop, _answer({
                "source": "parts", "columns": [{"field": "part_number"}]})
                if self.stop == "end_turn" else None)

    class FakeClient:
        stop = "end_turn"

        def __init__(self, **kw):
            captured.append({"client": kw})
            self.beta = SimpleNamespace(messages=FakeMessages(
                FakeClient.stop))

        async def close(self):
            pass

    real = ra.anthropic.AsyncAnthropic
    ra.anthropic.AsyncAnthropic = FakeClient
    try:
        stop, ans, model = asyncio.run(ra._call_model(
            "sk-x", "claude-opus-5-5", "Request: parts"))
        assert stop == "end_turn" and ans.spec.source == "parts"
        kw = captured[-1]
        assert kw["model"] == "claude-opus-5-5"
        assert kw["betas"] == ["server-side-fallback-2026-07-01"]
        assert kw["extra_body"] == {"fallbacks": "default"}
        assert kw["thinking"] == {"type": "adaptive"}
        assert kw["output_config"] == {"effort": "medium"}
        assert kw["output_format"] is ra.AssistantAnswer
        assert kw["system"][0]["cache_control"] == {"type": "ephemeral"}
        assert "CATALOG" in kw["system"][0]["text"]
        assert captured[-2]["client"]["api_key"] == "sk-x"
        stop, ans, _ = asyncio.run(ra._call_model(
            "sk-x", "claude-haiku-5-5", "Request: parts"))
        assert "betas" not in captured[-1] and \
            "extra_body" not in captured[-1]
        FakeClient.stop = "refusal"
        stop, ans, _ = asyncio.run(ra._call_model("k", "claude-opus-5-5",
                                                  "x"))
        assert stop == "refusal" and ans is None
        FakeClient.stop = "invalid"
        stop, ans, _ = asyncio.run(ra._call_model("k", "claude-opus-5-5",
                                                  "x"))
        assert stop == "invalid" and ans is None
    finally:
        ra.anthropic.AsyncAnthropic = real
    # Esquema de salida: el grupo de filtros NO es recursivo.
    sch_json = ra.AssistantAnswer.model_json_schema()
    grp = sch_json["$defs"]["LlmGroup"]["properties"]["items"]["items"]
    assert grp == {"$ref": "#/$defs/LlmCond"}
    spec = ra.llm_to_spec(_answer({
        "source": "work_orders", "columns": [{"field": "unit"}],
        "time": {"field": "created_at", "preset": "custom",
                 "date_from": "2026-01-01", "date_to": "2026-02-01"},
        "limit": 999999}).spec)
    assert spec["time"] == {"field": "created_at", "preset": "custom",
                            "from": "2026-01-01", "to": "2026-02-01"}
    assert spec["limit"] == eng.MAX_LIMIT
    print("OK asistente: request (fallbacks, caché, effort, esquema sin "
          "recursión)")


# ----- RBAC + rutas --------------------------------------------------------------------------

def test_rbac_and_route_order():
    assert _scope_for("POST", "/api/reports/run") is None
    assert _scope_for("POST", "/api/reports/export") is None
    assert _scope_for("POST", "/api/reports/assistant") is None
    assert _scope_for("POST", "/api/reports/saved") is None
    assert _scope_for("PATCH", "/api/reports/saved/3") is None
    assert _scope_for("DELETE", "/api/reports/saved/3") is None
    assert _scope_for("POST", "/api/reports/favorites") is None
    assert _scope_for("DELETE", "/api/reports/favorites") is None
    assert _scope_for("POST", "/api/reports/schedules") == "maint.edit"
    assert _scope_for("PATCH", "/api/reports/schedules/1") == "maint.edit"
    assert _scope_for("POST", "/api/reports/schedules/1/send-now") == \
        "maint.edit"
    assert _scope_for("POST", "/api/reports/cpm/refresh") == "maint.edit"
    assert _scope_for("GET", "/api/reports/library") is None
    assert _scope_for("POST", "/api/org") == "settings.manage"
    paths = [r.path for r in routes.router.routes
             if "GET" in getattr(r, "methods", set())]
    legacy = paths.index("/api/reports/{report_id}/download")
    for p in ("/api/reports/library", "/api/reports/sources",
              "/api/reports/schedules", "/api/reports/assistant/status",
              "/api/reports/system/{slug}", "/api/reports/saved/{report_id}"):
        assert paths.index(p) < legacy, p
    print("OK RBAC (run/export/assistant/saved/favoritos = autenticado; schedules = "
          "maint.edit) + orden de rutas")


def test_finisher_reports():
    """v2.22 finisher: copia de un guardado ajeno (based_on 'saved:<id>'),
    export sin fórmulas en encabezado/título y edición de un schedule de un
    reporte privado de otro usuario."""
    _seed()
    from app.core import report_export
    org_r = routes.reports_saved_create(routes.SavedReportIn(
        name="Org base", spec=SPEC, visibility="org"), MECH)
    copy = routes.reports_saved_create(routes.SavedReportIn(
        name="My copy", spec=SPEC, based_on=org_r["ref"]), VIEW)
    assert copy["based_on"] == org_r["ref"], copy["based_on"]
    priv = routes.reports_saved_create(routes.SavedReportIn(
        name="Private", spec=SPEC), VIEW)
    _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="Steal", spec=SPEC, based_on=priv["ref"]), MECH)
    _http(routes.reports_saved_create, 400, routes.SavedReportIn(
        name="Bad", spec=SPEC, based_on="saved:abc"), MECH)
    # Export: ninguna celda fórmula (datos, encabezados, título).
    res = {"columns": [{"key": "a", "label": '=HYPERLINK("http://x","c")',
                        "type": "text"},
                       {"key": "n", "label": "+n", "type": "number"}],
           "rows": [{"a": "=1+1", "n": -5}], "totals": {"n": -5}}
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(report_export.to_xlsx(res, "=cmd|x", "sub")))
    cells = [c for row in wb.active.iter_rows() for c in row
             if c.value is not None]
    assert cells and all(c.data_type != "f" for c in cells),         [(c.coordinate, c.data_type) for c in cells]
    assert wb.active["A1"].value == "=cmd|x"
    csv_text = report_export.to_csv(res).decode("utf-8-sig")
    head = csv_text.splitlines()[0]
    assert head.startswith("\"'=HYPERLINK") and ",'+n" in head, head
    assert "-5" in csv_text.splitlines()[1]          # números intactos
    # Schedule de un reporte 'mine' de MECH: DISP (maint.edit) edita
    # destinatarios/hora aunque no vea el reporte; cambiar a un reporte
    # ajeno sigue validando.
    mine = routes.reports_saved_create(routes.SavedReportIn(
        name="Mech private", spec=SPEC), MECH)
    sch = routes.reports_schedules_create(routes.ReportScheduleIn(
        report_ref=mine["ref"], recipients="a@example.com"), MECH)
    up = routes.reports_schedules_patch(sch["id"], routes.ReportSchedulePatch(
        report_ref=mine["ref"], version="", recipients="b@example.com",
        hour=9), DISP)
    assert up["recipients"] == ["b@example.com"] and up["hour"] == 9, up
    _http(routes.reports_schedules_patch, 400, sch["id"],
          routes.ReportSchedulePatch(report_ref=priv["ref"]), DISP)
    routes.reports_schedules_delete(sch["id"])
    # Asistente: un group_by con una dimensión que no es columna se vuelve
    # columna (el builder agrupa por columnas: preview == spec validada).
    a = _answer({"source": "wo_lines",
                 "columns": [{"field": "unit"},
                             {"field": "amount", "agg": "sum"}],
                 "group_by": ["unit", "kind"]})
    out = ra.llm_to_spec(a.spec)
    assert [c["field"] for c in out["columns"]] == ["unit", "kind", "amount"]
    assert out["group_by"] == ["unit", "kind"]
    print("OK finisher: based_on saved:<id>, export sin fórmulas, edición "
          "de schedule ajeno, group_by del asistente")


if __name__ == "__main__":
    test_saved_reports_visibility_and_permissions()
    test_run_route_pii_and_export()
    test_next_run_math()
    test_schedules_crud_claim_and_send()
    test_scheduler_iterates_orgs()
    test_mailer_attachments()
    test_assistant_unavailable_and_settings()
    test_assistant_flow_with_fake_model()
    test_assistant_request_shape()
    test_rbac_and_route_order()
    test_finisher_reports()
    print("\nALL REPORT SAVED/SCHEDULE/ASSISTANT TESTS PASSED")
