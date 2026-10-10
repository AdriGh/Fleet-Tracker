# -*- coding: utf-8 -*-
"""Tests del motor de reportes v2.22 (registro, specs, biblioteca, paridad).

`python backend/tests/test_reports_engine.py`. DB SQLite temporal fresca;
la flota viva (Samsara), los tableros de PM y el roster de conductores se
reemplazan por fakes (sin red).

Lo que NO hay que perder:
  - registro: toda fuente SQL corre con TODOS sus campos; catálogo para el
    builder; campos ocultos no se aceptan;
  - validación: fuente/campo/op/agg desconocidos o mal tipados -> ValueError
    con mensaje; group_by vs agregados; tiempo solo sobre campos de fecha
    (el 'period' en las fuentes de período); custom con fechas válidas;
    límite 1..50,000; PII no filtrable/ordenable/agrupable sin pii.view;
  - presets relativos a la fecha de la ORG;
  - filtros: misma semántica empujados a SQL o evaluados en Python; grupos
    any/all; texto case-insensitive; datetimes por fecha de la org;
  - agregación (count, count_distinct, sum, avg, min, max), orden con
    vacíos al final, totales, límite + truncated + aviso;
  - tenant: sin org en contexto NO corre; otra org no ve nada (filtro
    explícito por join);
  - PII enmascarada en customers sin pii.view;
  - biblioteca: >= 30 reportes, TODAS las versiones validan y corren;
  - PARIDAD: spend_lines == spend_report y cpm == cpm_report;
  - export CSV (BOM, fórmulas neutralizadas) y XLSX (strings, no fórmulas).
"""
import asyncio
import io
import os
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

_BACKEND = Path(__file__).resolve().parent.parent
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

_TMP = Path(tempfile.mkdtemp())
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP / 'reports_engine_test.db'}"

from app import db                                              # noqa: E402
from app.core import (                                          # noqa: E402
    customers, fleet_inventory, maint, odometer, org_config, parts,
    purchasing, report_engine as eng, report_export, report_library as lib,
    report_registry as reg, reports, technicians, tenant, terminals, tms,
    workorders,
)

ORG = db.default_org_id()
tenant.set_current_org(ORG)


# ----- Fakes (sin red) ---------------------------------------------------------

FLEET = [
    {"unit": "MEM 77", "unit_type": "truck", "make": "Freightliner",
     "model": "Cascadia", "year": "2021", "company": "CHASER",
     "source": "samsara", "open_defects": 1},
    {"unit": "TRL 500", "unit_type": "trailer", "make": "Wabash",
     "model": "DuraPlate", "year": "2019", "company": "CHASER",
     "source": "samsara", "open_defects": 0},
    {"unit": "DAL 12", "unit_type": "truck", "make": "International",
     "model": "LT", "year": "2020", "company": "MCC", "source": "samsara",
     "open_defects": 0},
]


async def _fake_fleet(refresh: bool = False) -> dict:
    return {"available": True, "source": "fake",
            "units": [dict(u) for u in FLEET]}


async def _fake_board(kind: str, refresh: bool = False) -> dict:
    if kind == "pm":
        units = [
            {"unit": "MEM 77", "model": "2021 Freightliner Cascadia",
             "driver": "Ana", "current_miles": 120000, "last_miles": 95000,
             "last_date": "2026-03-01", "next_due_miles": 115000,
             "to_due": -5000, "status": "overdue", "due_status": "overdue",
             "ops_status": "", "due_by": "miles"},
            {"unit": "DAL 12", "model": "2020 International LT",
             "driver": "", "current_miles": 50000, "last_miles": 40000,
             "last_date": "2026-08-01", "next_due_miles": 60000,
             "to_due": 10000, "status": "on_track",
             "due_status": "on_track", "ops_status": "", "due_by": "miles"},
        ]
        return {"kind": "pm", "available": True, "units": units,
                "upcoming_miles": 5500,
                "interval_miles": 20000, "errors": []}
    if kind == "dot":
        units = [
            {"unit": "TRL 500", "model": "", "driver": "",
             "last_date": "2025-10-20", "next_due_date": "2026-10-20",
             "to_due": 11, "status": "upcoming", "due_status": "upcoming",
             "ops_status": "", "due_by": "days", "current_miles": None,
             "last_miles": None},
        ]
        return {"kind": "dot", "available": True, "units": units,
                "upcoming_days": 30,
                "interval_days": 365, "errors": []}
    units = [
        {"unit": "TRL 500", "model": "Thermo King", "current_hours": 3200.0,
         "current_hours_date": "2026-10-01", "last_date": "2026-01-10",
         "last_hours": 1600.0, "next_due_hours": 3100.0,
         "next_due_date": "2027-01-10", "to_due": -100.0, "to_due_days": 93,
         "due_by": "hours", "status": "overdue", "due_status": "overdue",
         "ops_status": "", "meter_missing": False, "current_miles": None},
    ]
    return {"kind": "reefer", "available": True, "units": units,
            "upcoming_hours": 150,
            "upcoming_days": 30, "interval_hours": 1500, "errors": []}


DRIVERS = [
    {"name": "Ana Lopez", "company": "CHASER", "phone": "901-555-1234",
     "email": "ana@example.com", "license_number": "D123",
     "license_state": "TN", "has_profile": True, "role": "owner_operator",
     "truck": "MEM 77", "trailer": "TRL 500", "hired_date": "2024-01-02",
     "cdl_exp": "2026-10-01", "med_exp": "2027-05-01", "mvr_exp": "",
     "chouse_exp": "2026-10-20"},
    {"name": "Bob Ray", "company": "MCC", "phone": "", "email": "",
     "license_number": "", "license_state": "", "has_profile": False,
     "role": "owner_operator", "truck": "", "trailer": "",
     "hired_date": "", "cdl_exp": "", "med_exp": "", "mvr_exp": "",
     "chouse_exp": ""},
]


async def _fake_drivers() -> list[dict]:
    return [dict(d) for d in DRIVERS]


fleet_inventory.load = _fake_fleet
maint.board = _fake_board
tms.list_drivers = _fake_drivers


# ----- Helpers -------------------------------------------------------------------

def _run(spec, mask_pii=False):
    return asyncio.run(eng.run(spec, mask_pii=mask_pii))


def _raises(fn, needle: str, *args, **kw):
    try:
        res = fn(*args, **kw)
        if asyncio.iscoroutine(res):
            asyncio.run(res)
    except ValueError as exc:
        assert needle in str(exc), f"{needle!r} not in {exc!s}"
        return
    raise AssertionError(f"debió lanzar ValueError con {needle!r}")


def _close(a, b, tol=0.011):
    assert abs((a or 0) - (b or 0)) <= tol, (a, b)


def _wo(unit, service_date, lines, *, company="", campaign="", is_pm=False,
        status="open", mechanic="Luis", wo_type="inside", vendor_id=None,
        customer_id=0):
    w = workorders.create_wo(unit, f"svc {unit}", company=company,
                             service_date=service_date, campaign=campaign,
                             is_pm=is_pm, mechanic=mechanic, wo_type=wo_type,
                             vendor_id=vendor_id, customer_id=customer_id)
    for kind, desc, qty, cost, pn in lines:
        workorders.add_line(w["id"], kind, desc, qty, cost, pn)
    if status != "open":
        workorders.update_wo(w["id"], {"status": status})
    return w["id"]


SEEDED = {}


def _raw_line(wo_id, kind, desc, qty, cost):
    with db.SessionLocal() as s:
        s.add(db.WorkOrderLine(wo_id=wo_id, kind=kind, description=desc,
                               qty=qty, unit_cost=cost, part_number="",
                               org_id=ORG))
        s.commit()


def _seed():
    if SEEDED:
        return
    terminals.save_terminal("", "Memphis", ["MEM", "TRL"])
    terminals.save_terminal("", "Dallas", ["DAL"])
    v = parts.create_vendor({"name": "Speedco", "kind": "shop",
                             "city": "Memphis", "state": "TN"})
    parts.create_part({"part_number": "BP-1", "description": "Brake pad",
                       "category": "brakes", "cost": 40, "on_hand": 5,
                       "reorder_point": 6})
    parts.create_part({"part_number": "OF-9", "description": "Oil filter",
                       "category": "", "cost": 12, "on_hand": 20})
    cust = customers.create_customer({"name": "Juan Perez",
                                      "email": "juan@example.com",
                                      "phone": "901-555-7777",
                                      "parts_markup_pct": 10,
                                      "labor_markup_pct": 20})
    SEEDED["customer"] = cust["id"]
    SEEDED["vendor"] = v["id"]
    # Gasto (fechas fijas en 2026).
    SEEDED["w1"] = _wo("MEM 77", "2026-06-01",
                       [("part", "brake pad", 2, 40, "BP-1"),
                        ("labor", "brake labor", 3, 100, "")],
                       company="CHASER", status="in_progress")
    SEEDED["w2"] = _wo("MEM 77", "2026-06-15",
                       [("part", "oil filter", 1, 12, "OF-9"),
                        ("part", "engine oil 15w40", 10, 9.5, ""),
                        ("labor", "pm labor", 1.5, 100, "")],
                       company="CHASER", campaign="pm", status="completed")
    SEEDED["w3"] = _wo("DAL 12", "2026-07-03",
                       [("part", "295/75 tire", 2, 410, "")],
                       company="MCC", status="open")
    _raw_line(SEEDED["w3"], "part", "core charge credit", 1, -50)
    SEEDED["w4"] = _wo("DAL 12", "2026-07-20",
                       [("part", "starter", 1, 300, "")], company="MCC")
    # Crédito directo en la base (la API no deja costos negativos, pero
    # spend_report los admite): la WO queda con neto < 0 -> cuenta en el
    # total pero NO en wo_count ni en el split preventivo/reactivo.
    _raw_line(SEEDED["w4"], "part", "credit", 1, -400)
    SEEDED["w5"] = _wo("TRL 500", "2026-08-02",
                       [("labor", "dot inspection", 1, 150, "")],
                       campaign="dot", status="completed")
    SEEDED["w6"] = _wo("MEM 77", "2026-09-10", [], status="open")  # vacía
    SEEDED["w7"] = _wo("TRL 500", "2026-09-12",
                       [("labor", "outside reefer repair", 2, 125, "")],
                       wo_type="outside", vendor_id=v["id"], mechanic="")
    # Customer: facturada (precio congelado con markup).
    SEEDED["w8"] = _wo("MEM 77", "2026-09-20",
                       [("part", "brake pad", 1, 40, "BP-1"),
                        ("labor", "install", 1, 100, "")],
                       customer_id=cust["id"])
    workorders.update_wo(SEEDED["w8"], {"status": "invoiced"})
    # Odómetro para el CPM.
    for unit, pts in {"MEM 77": [("2026-06-01", 100000),
                                 ("2026-06-30", 104000),
                                 ("2026-07-31", 109000)],
                      "DAL 12": [("2026-07-01", 50000),
                                 ("2026-07-31", 52500)]}.items():
        for d, mi in pts:
            odometer.log_reading(unit, d, mi)
    # PO recibida + ajuste.
    po = purchasing.create_po("Speedco", lines=[
        {"part_number": "BP-1", "description": "Brake pad", "qty": 4,
         "unit_cost": 38}])
    purchasing.update_po(po["id"], {"status": "ordered"}, actor="Ana")
    line_id = po["lines"][0]["id"] if po.get("lines") else \
        purchasing.get_po(po["id"])["lines"][0]["id"]
    purchasing.receive_po(po["id"], [{"line_id": line_id, "qty_now": 4}],
                          actor="Ana")
    purchasing.add_adjustment(po["id"], "payment", 50, actor="Ana")
    SEEDED["po"] = po["id"]
    technicians.ensure("Luis")


# ----- Tests ---------------------------------------------------------------------------

def test_registry_and_catalog():
    _seed()
    cat = reg.catalog()
    keys = {s["key"] for s in cat}
    for must in ("work_orders", "wo_lines", "parts", "stock_movements",
                 "purchase_orders", "po_lines", "po_adjustments", "vendors",
                 "service_requests", "maint_records", "odometer_readings",
                 "meter_readings", "defects", "customers", "units",
                 "technicians", "pm_due", "warranty_claims", "core_items",
                 "spend_lines", "cpm"):
        assert must in keys, must
    for s in cat:
        assert s["description"] and s["fields"], s["key"]
        assert all(f["type"] in reg.TYPES for f in s["fields"])
        assert not any(f["name"].startswith("cust_") for f in s["fields"])
    # Todas las fuentes SQL corren con todos sus campos (2 tandas de 40).
    for src in reg.SOURCES.values():
        if not src.is_sql:
            continue
        names = [n for n, f in src.fields.items() if not f.hidden]
        for i in range(0, len(names), 40):
            r = _run({"source": src.key,
                      "columns": [{"field": n} for n in names[i:i + 40]],
                      "totals": True})
            assert r["source"] == src.key and "rows" in r
    # Campo oculto: no se acepta.
    _raises(_run, "unknown field 'cust_parts_pct'",
            {"source": "work_orders", "columns": [{"field":
                                                    "cust_parts_pct"}]})
    print("OK registro: catálogo + todas las fuentes SQL corren")


def test_validation_errors():
    _raises(_run, "unknown source", {"source": "nope",
                                     "columns": [{"field": "x"}]})
    _raises(_run, "unknown field 'bogus'",
            {"source": "work_orders", "columns": [{"field": "bogus"}]})
    _raises(_run, "invalid report spec",
            {"source": "work_orders", "columns": []})
    _raises(_run, "invalid report spec",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "filters": {"items": [{"field": "unit", "op": "like",
                                    "value": "x"}]}})
    _raises(_run, "invalid report spec",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "colour": "red"})
    _raises(_run, "'sum' needs a number",
            {"source": "work_orders", "columns": [{"field": "unit",
                                                   "agg": "sum"}]})
    _raises(_run, "must be in group_by",
            {"source": "work_orders",
             "columns": [{"field": "unit"}, {"field": "total", "agg": "sum"}],
             "group_by": ["status"]})
    _raises(_run, "group_by needs at least one column",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "group_by": ["unit"]})
    _raises(_run, "'contains' only works on text",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "filters": {"items": [{"field": "total", "op": "contains",
                                    "value": "1"}]}})
    _raises(_run, "needs a number",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "filters": {"items": [{"field": "total", "op": "gt",
                                    "value": "lots"}]}})
    _raises(_run, "needs two values",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "filters": {"items": [{"field": "total", "op": "between",
                                    "value": [1]}]}})
    _raises(_run, "time field 'unit'",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "time": {"field": "unit", "preset": "mtd"}})
    _raises(_run, "period field",
            {"source": "cpm", "columns": [{"field": "unit"}],
             "time": {"field": "unit", "preset": "mtd"}})
    _raises(_run, "custom time range needs",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "time": {"field": "created_at", "preset": "custom"}})
    _raises(_run, "'from' is after 'to'",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "time": {"field": "created_at", "preset": "custom",
                      "from": "2026-05-01", "to": "2026-04-01"}})
    _raises(_run, "invalid report spec",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "limit": 50001})
    _raises(_run, "appears twice",
            {"source": "work_orders",
             "columns": [{"field": "unit"}, {"field": "unit"}]})
    _raises(_run, "sort field 'title' must be one of the report columns",
            {"source": "work_orders",
             "columns": [{"field": "unit"}, {"field": "id", "agg": "count"}],
             "sort": [{"field": "title"}]})
    _raises(_run, "chart field",
            {"source": "work_orders", "columns": [{"field": "unit"}],
             "chart": {"type": "bar", "x": "status", "y": []}})
    # PII: sin pii.view no se filtra/ordena/agrupa por contacto.
    for spec in (
            {"source": "customers", "columns": [{"field": "name"}],
             "filters": {"items": [{"field": "email", "op": "contains",
                                    "value": "x"}]}},
            {"source": "customers", "columns": [{"field": "name"}],
             "sort": [{"field": "phone"}]},
            {"source": "customers", "columns": [{"field": "email"},
                                                {"field": "id",
                                                 "agg": "count"}]}):
        _raises(_run, "permission to see contact info", spec, mask_pii=True)
    print("OK validación: errores claros en inglés")


def test_presets():
    t = date(2026, 10, 9)                      # viernes
    assert eng.preset_range("today", t) == (t, t)
    assert eng.preset_range("yesterday", t) == (date(2026, 10, 8),) * 2
    assert eng.preset_range("this_week", t) == (date(2026, 10, 5),
                                                date(2026, 10, 11))
    assert eng.preset_range("prev_week", t) == (date(2026, 9, 28),
                                                date(2026, 10, 4))
    assert eng.preset_range("mtd", t) == (date(2026, 10, 1), t)
    assert eng.preset_range("prev_month", t) == (date(2026, 9, 1),
                                                 date(2026, 9, 30))
    assert eng.preset_range("last_30d", t) == (date(2026, 9, 10), t)
    assert eng.preset_range("last_90d", t) == (date(2026, 7, 12), t)
    assert eng.preset_range("qtd", t) == (date(2026, 10, 1), t)
    assert eng.preset_range("ytd", t) == (date(2026, 1, 1), t)
    assert eng.preset_range("last_12m", t) == (date(2025, 10, 10), t)
    assert eng.preset_range("last_12m", date(2028, 2, 29)) == (
        date(2027, 3, 1), date(2028, 2, 29))
    assert eng.preset_range("all", t) == (None, None)
    assert eng.preset_range("custom", t, "2026-01-01", "") == (
        date(2026, 1, 1), None)
    print("OK presets relativos a la fecha de la org")


def test_filters_aggregation_sort_totals():
    _seed()
    base = {"source": "wo_lines",
            "columns": [{"field": "wo_id"}, {"field": "unit"},
                        {"field": "amount"}, {"field": "kind"}]}
    allr = _run(base)["rows"]
    assert len(allr) == 13, len(allr)
    # eq / ne / in / not_in / gt / between / contains / empty (texto
    # case-insensitive; SQL empujado y Python dan lo mismo).
    def n(items, match="all"):
        return len(_run({**base, "filters": {"match": match,
                                             "items": items}})["rows"])
    assert n([{"field": "kind", "op": "eq", "value": "LABOR"}]) == 5
    assert n([{"field": "unit", "op": "eq", "value": "mem 77"}]) == 7
    assert n([{"field": "kind", "op": "ne", "value": "labor"}]) == 8
    assert n([{"field": "unit", "op": "in",
               "value": ["DAL 12", "TRL 500"]}]) == 6
    assert n([{"field": "unit", "op": "not_in", "value": "DAL 12"}]) == 9
    assert n([{"field": "amount", "op": "gt", "value": 300}]) == 1
    assert n([{"field": "amount", "op": "between", "value": [95, 300]}]) == 7
    assert n([{"field": "description", "op": "contains",
               "value": "BRAKE"}]) == 3
    assert n([{"field": "part_number", "op": "empty"}]) == 10
    assert n([{"field": "part_number", "op": "not_empty"}]) == 3
    assert n([{"field": "qty", "op": "gte", "value": 2}]) == 5  # pushed SQL
    assert n([{"field": "amount", "op": "lt", "value": 0},
              {"field": "unit", "op": "eq", "value": "MEM 77"}],
             "any") == 9
    assert n([{"field": "unit", "op": "eq", "value": "DAL 12"},
              {"match": "any", "items": [
                  {"field": "amount", "op": "lt", "value": 0},
                  {"field": "description", "op": "contains",
                   "value": "tire"}]}]) == 3
    # Agregación + orden + totales.
    r = _run({"source": "wo_lines",
              "columns": [{"field": "unit"},
                          {"field": "amount", "agg": "sum"},
                          {"field": "id", "agg": "count"},
                          {"field": "wo_id", "agg": "count_distinct"},
                          {"field": "amount", "agg": "avg"},
                          {"field": "amount", "agg": "max"},
                          {"field": "date", "agg": "min"}],
              "sort": [{"field": "sum_amount", "dir": "desc"}],
              "totals": True})
    units = [x["unit"] for x in r["rows"]]
    assert units == ["MEM 77", "DAL 12", "TRL 500"], units
    mem = r["rows"][0]
    assert mem["sum_amount"] == 2 * 40 + 300 + 12 + 95 + 150 + 40 + 100
    assert mem["count_id"] == 7 and mem["count_distinct_wo_id"] == 3
    assert mem["min_date"] == "2026-06-01"
    assert r["totals"]["count_id"] == 13
    assert r["totals"]["count_distinct_wo_id"] == 7
    _close(r["totals"]["sum_amount"], sum(x["sum_amount"] for x in r["rows"]))
    assert {c["key"]: c["type"] for c in r["columns"]}["count_id"] == "number"
    # Orden: vacíos al final en asc y desc.
    for d in ("asc", "desc"):
        rows = _run({"source": "work_orders",
                     "columns": [{"field": "display_no"},
                                 {"field": "vendor_name"}],
                     "sort": [{"field": "vendor_name", "dir": d}]})["rows"]
        assert rows[0]["vendor_name"] == "Speedco"
        assert rows[-1]["vendor_name"] is None
    # Límite + truncated + aviso "sin filtro de fecha".
    r = _run({**base, "limit": 5})
    assert r["row_count"] == 5 and r["total_count"] == 13 and r["truncated"]
    assert any("No date filter" in w for w in r["warnings"])
    # Totales en detalle: suma columnas monetarias, no ids.
    r = _run({**base, "totals": True})
    assert "amount" in r["totals"] and "wo_id" not in r["totals"]
    print("OK filtros (SQL = Python), agregados, orden, totales, límite")


def test_text_pushdown_autoescape():
    """Texto: case-insensitive empujado a SQL (ASCII) con LIKE autoescape
    (% y _ del usuario son literales); no-ASCII se evalúa en Python."""
    _seed()
    ids = [workorders.create_wo("ESC 1", "50% off brake_pad")["id"],
           workorders.create_wo("ESC 2", "500 off brakeXpad")["id"],
           workorders.create_wo("ESC 3", "Ñandú repair")["id"]]

    def n(cond):
        return len(_run({"source": "work_orders",
                         "columns": [{"field": "title"}],
                         "filters": {"items": [
                             {"field": "unit", "op": "contains",
                              "value": "esc "}, cond]}})["rows"])
    assert n({"field": "title", "op": "contains", "value": "50%"}) == 1
    assert n({"field": "title", "op": "contains", "value": "_pad"}) == 1
    assert n({"field": "title", "op": "contains", "value": "ñandú"}) == 1
    assert n({"field": "unit", "op": "eq", "value": "esc 1"}) == 1
    assert n({"field": "unit", "op": "not_in",
              "value": ["ESC 1", "esc 3"]}) == 1
    assert n({"field": "title", "op": "ne",
              "value": "500 OFF BRAKEXPAD"}) == 2
    with db.SessionLocal() as s:
        for i in ids:
            s.delete(s.get(db.WorkOrder, i))
        s.commit()
    print("OK texto: SQL case-insensitive + LIKE con autoescape")


def test_time_filters_and_derived():
    _seed()
    r = _run({"source": "wo_lines",
              "columns": [{"field": "amount", "agg": "sum"}],
              "time": {"field": "date", "preset": "custom",
                       "from": "2026-06-01", "to": "2026-06-30"}})
    assert r["rows"][0]["sum_amount"] == 80 + 300 + 12 + 95 + 150
    assert r["time"]["from"] == "2026-06-01"
    # created_at (datetime): hoy en la zona de la org -> todas.
    r = _run({"source": "work_orders", "columns": [{"field": "id",
                                                    "agg": "count"}],
              "time": {"field": "created_at", "preset": "today"}})
    assert r["rows"][0]["count_id"] == 8
    r = _run({"source": "work_orders", "columns": [{"field": "id",
                                                    "agg": "count"}],
              "time": {"field": "created_at", "preset": "yesterday"}})
    assert r["rows"] == [{"count_id": 0}]      # KPI sin filas = una fila 0
    # Derivados de la WO.
    rows = {x["display_no"]: x for x in _run({
        "source": "work_orders",
        "columns": [{"field": "display_no"}, {"field": "total"},
                    {"field": "parts_total"}, {"field": "labor_hours"},
                    {"field": "billable_total"}, {"field": "bill_to"},
                    {"field": "terminal"}, {"field": "spend_type"},
                    {"field": "effective_date"}, {"field": "days_open"},
                    {"field": "balance"}, {"field": "vendor_name"}]})["rows"]}
    w1 = rows[str(SEEDED["w1"])]
    assert w1["total"] == 380 and w1["parts_total"] == 80
    assert w1["labor_hours"] == 3 and w1["terminal"] == "Memphis"
    assert w1["bill_to"] == "Self" and w1["spend_type"] == "Reactive"
    assert w1["effective_date"] == "2026-06-01" and w1["days_open"] == 0
    w8 = rows[str(SEEDED["w8"])]
    assert w8["bill_to"] == "Juan Perez"
    assert w8["billable_total"] == 44 + 120 and w8["balance"] == 164
    assert rows[str(SEEDED["w2"])]["spend_type"] == "Preventive"
    assert rows[str(SEEDED["w7"])]["vendor_name"] == "Speedco"
    assert rows[str(SEEDED["w4"])]["terminal"] == "Dallas"
    # Partes: valuación, bajo punto de pedido, uso.
    pr = {x["part_number"]: x for x in _run({
        "source": "parts",
        "columns": [{"field": "part_number"}, {"field": "total_value"},
                    {"field": "below_reorder"}, {"field": "qty_used"},
                    {"field": "last_used"}, {"field": "on_hand"}]})["rows"]}
    assert pr["BP-1"]["below_reorder"] is False and pr["BP-1"]["qty_used"] == 3
    assert pr["BP-1"]["last_used"] == org_config.org_today().isoformat()
    assert pr["OF-9"]["total_value"] == 240
    # PO: pagado / balance desde ajustes.
    po = _run({"source": "purchase_orders",
               "columns": [{"field": "id"}, {"field": "total"},
                           {"field": "paid"}, {"field": "balance"},
                           {"field": "line_count"}, {"field": "status"}]})
    p = po["rows"][0]
    assert p["total"] == 152 and p["paid"] == 50 and p["balance"] == 102
    assert p["line_count"] == 1 and p["status"] == "received"
    print("OK rango de tiempo (fecha y datetime en la org) + derivados")


def test_pii_masking_and_tenant_isolation():
    _seed()
    spec = {"source": "customers",
            "columns": [{"field": "name"}, {"field": "email"},
                        {"field": "phone"}, {"field": "billed_all"},
                        {"field": "open_balance"}]}
    clear = _run(spec)["rows"][0]
    assert clear["email"] == "juan@example.com"
    assert clear["billed_all"] == 164 and clear["open_balance"] == 164
    masked = _run(spec, mask_pii=True)
    assert "•••" in masked["rows"][0]["email"]
    assert masked["rows"][0]["phone"].endswith("7777")
    assert masked["pii_masked"] is True
    # Otra org: no ve nada; sin tenant no corre.
    with db.SessionLocal() as s:
        o = db.Organization(slug="other-rep", name="Other", active=True,
                            created_at=datetime.now())
        s.add(o)
        s.commit()
        other = o.id
    tok = tenant.set_current_org(other)
    try:
        r = _run({"source": "wo_lines", "columns": [{"field": "id"}]})
        assert r["rows"] == []
        r = _run({"source": "work_orders",
                  "columns": [{"field": "id"}, {"field": "vendor_name"},
                              {"field": "total"}]})
        assert r["rows"] == []
    finally:
        tenant.reset_current_org(tok)
    tok = tenant.set_current_org(None)
    try:
        _raises(_run, "no organization in context",
                {"source": "work_orders", "columns": [{"field": "id"}]})
    finally:
        tenant.reset_current_org(tok)
    print("OK PII enmascarada sin pii.view + aislamiento por tenant")


def test_library_validates_and_runs():
    _seed()
    reps = lib.list_system()
    assert len(reps) >= 30, len(reps)
    slugs = [r["slug"] for r in reps]
    assert len(slugs) == len(set(slugs))
    assert set(r["category"] for r in reps) <= set(lib.CATEGORIES)
    sr = [r for r in reps if r["origin"] == "squarerigger"]
    assert len(sr) >= 25, len(sr)
    for r in reps:
        assert r["versions"] and r["default_version_label"], r["slug"]
        for v in r["versions"]:
            spec, ver = lib.spec_for(r["slug"], v["key"])
            eng.validate(spec)
            res = _run(spec)
            assert "rows" in res, (r["slug"], v["key"])
            assert ver["key"] == v["key"]
    _raises(lib.spec_for, "unknown version", "cost-per-mile", "nope")
    _raises(lib.spec_for, "unknown system report", "nope")
    # Spot checks sobre los providers (fakes).
    due = _run(lib.spec_for("pms-due-and-overdue")[0])["rows"]
    assert {(d["unit"], d["status"]) for d in due} == {
        ("MEM 77", "overdue"), ("TRL 500", "overdue"),
        ("TRL 500", "upcoming")}
    reefer = _run(lib.spec_for("reefer-pm-due")[0])["rows"]
    assert reefer[0]["unit"] == "TRL 500" and reefer[0]["to_due"] == -100
    assert reefer[0]["unit_label"] == "hrs"
    k20 = _run(lib.spec_for("pm-due-next-20k-miles")[0])["rows"]
    assert [x["unit"] for x in k20] == ["MEM 77", "DAL 12"]
    assert k20[0]["overdue_by"] == 5000
    since = _run(lib.spec_for("mileage-since-last-service")[0])["rows"]
    assert since[0]["since_last"] == 25000
    docs = _run(lib.spec_for("driver-compliance")[0])["rows"]
    assert any(d["driver"] == "Bob Ray" and d["state"] == "missing"
               for d in docs)
    comp = _run(lib.spec_for("pm-compliance")[0])
    assert comp["totals"]["count_unit"] == 3
    snap = _run(lib.spec_for("fleet-mileage-snapshot")[0])["rows"]
    mem = next(x for x in snap if x["unit"] == "MEM 77")
    assert mem["odometer"] == 109000 and mem["odometer_date"] == "2026-07-31"
    cust = _run(lib.spec_for("equipment-by-customer")[0])["rows"]
    assert cust == []                          # sin dueños asignados
    customers.set_unit(SEEDED["customer"], "MEM 77")
    cust = _run(lib.spec_for("equipment-by-customer")[0])["rows"]
    assert [c["customer"] for c in cust] == ["Juan Perez"]
    bill = _run(lib.spec_for("customer-billing", "all")[0])
    assert bill["rows"] == [{"customer_name": "Juan Perez", "count_id": 1,
                             "sum_billable_total": 164.0,
                             "sum_billed_paid": 0.0,
                             "sum_balance": 164.0}], bill["rows"]
    lead = _run(lib.spec_for("vendor-lead-time", "ytd")[0])["rows"]
    assert lead and lead[0]["vendor"] == "Speedco"
    assert lead[0]["avg_lead_time_days"] == 0
    tech = _run(lib.spec_for("technician-leaderboard", "ytd")[0])["rows"]
    assert tech[0]["mechanic"] == "Luis"
    print(f"OK biblioteca: {len(reps)} reportes, todas las versiones "
          "validan y corren")


def _engine_spend(time, extra_filters=()):
    def q(columns, group_by, filters=()):
        spec = {"source": "spend_lines", "columns": columns,
                "group_by": group_by,
                "filters": {"items": list(extra_filters) + list(filters)},
                "limit": 50000}
        if time:
            spec["time"] = time
        return _run(spec)["rows"]
    return q


def test_spend_parity():
    _seed()
    cases = [("", "", None),
             ("2026-06-01", "2026-06-30",
              {"field": "date", "preset": "custom", "from": "2026-06-01",
               "to": "2026-06-30"}),
             ("2026-07-01", "2026-12-31",
              {"field": "date", "preset": "custom", "from": "2026-07-01",
               "to": "2026-12-31"})]
    for term in ("", "MEMPHIS"):
        cfg = terminals.get_all()
        key = next((t["key"] for t in cfg["terminals"]
                    if t["label"] == "Memphis"), "")
        term_key = key if term else ""
        filt = ([{"field": "terminal", "op": "eq", "value": term_key}]
                if term_key else [])
        for d_from, d_to, time in cases:
            rep = reports.spend_report(d_from, d_to, terminal=term_key,
                                       top_units=100, top_parts=100)
            q = _engine_spend(time, filt)
            t = rep["totals"]
            tot = q([{"field": "amount", "agg": "sum"}], [])
            _close(tot[0]["sum_amount"] if tot else 0, t["total_spend"])
            kinds = {r["kind"]: r["sum_amount"] for r in q(
                [{"field": "kind"}, {"field": "amount", "agg": "sum"}],
                ["kind"])}
            _close(kinds.get("labor"), t["labor_spend"])
            _close(kinds.get("part"), t["parts_spend"])
            types = {r["spend_type"]: r["sum_amount"] for r in q(
                [{"field": "spend_type"}, {"field": "amount", "agg": "sum"}],
                ["spend_type"])}
            _close(types.get("Preventive"), t["pm_spend"])
            _close(types.get("Reactive"), t["reactive_spend"])
            wo = q([{"field": "wo_id", "agg": "count_distinct"}], [],
                   [{"field": "spend_type", "op": "not_empty"}])
            assert (wo[0]["count_distinct_wo_id"] if wo else 0) == \
                t["wo_count"], (wo, t)
            cats = {r["category_label"]: r["sum_amount"] for r in q(
                [{"field": "category_label"},
                 {"field": "amount", "agg": "sum"}], ["category_label"])}
            assert set(cats) == {c["label"] for c in rep["by_category"]}
            for c in rep["by_category"]:
                _close(cats[c["label"]], c["value"])
            units = {r["unit"]: r["sum_amount"] for r in q(
                [{"field": "unit"}, {"field": "amount", "agg": "sum"}],
                ["unit"])}
            assert set(units) == {u["label"] for u in rep["by_unit"]}
            for u in rep["by_unit"]:
                _close(units[u["label"]], u["value"])
            months = {r["month"]: r["sum_amount"] for r in q(
                [{"field": "month"}, {"field": "amount", "agg": "sum"}],
                ["month"])}
            assert sorted(months) == [m["label"] for m in rep["by_month"]]
            for m in rep["by_month"]:
                _close(months[m["label"]], m["value"])
            tops = {r["part_number"]: r for r in q(
                [{"field": "part_number"}, {"field": "amount", "agg": "sum"},
                 {"field": "qty", "agg": "sum"}], ["part_number"],
                [{"field": "kind", "op": "eq", "value": "part"},
                 {"field": "part_number", "op": "not_empty"}])}
            assert set(tops) == {p["part_number"] for p in rep["top_parts"]}
            for p in rep["top_parts"]:
                _close(tops[p["part_number"]]["sum_amount"], p["value"])
                _close(tops[p["part_number"]]["sum_qty"], p["qty"])
    # La versión por defecto de "Spend Overview" usa la misma fuente.
    spec, _v = lib.spec_for("spend-overview", "by_unit")
    assert spec["source"] == "spend_lines"
    print("OK paridad: spend_lines == spend_report (totales, categorías, "
          "unidades, meses, top partes, PM/reactivo, terminal)")


def test_cpm_parity():
    _seed()
    for d_from, d_to, time in (
            ("", "", {"field": "period", "preset": "all"}),
            ("2026-06-01", "2026-06-30",
             {"field": "period", "preset": "custom", "from": "2026-06-01",
              "to": "2026-06-30"}),
            ("2026-06-01", "2026-07-31",
             {"field": "period", "preset": "custom", "from": "2026-06-01",
              "to": "2026-07-31"})):
        rep = reports.cpm_report(d_from, d_to)
        res = _run({"source": "cpm",
                    "columns": [{"field": "unit"}, {"field": "spend"},
                                {"field": "miles"}, {"field": "cpm"},
                                {"field": "labor_spend"},
                                {"field": "parts_spend"},
                                {"field": "labor_cpm"},
                                {"field": "parts_cpm"}, {"field": "period"}],
                    "time": time, "totals": True})
        got = {r["unit"]: r for r in res["rows"]}
        assert set(got) == {r["unit"] for r in rep["by_unit"]}
        for r in rep["by_unit"]:
            g = got[r["unit"]]
            _close(g["spend"], r["spend"])
            assert g["miles"] == r["miles"], (g, r)
            if r["cpm"] is None:
                assert g["cpm"] is None
            else:
                _close(g["cpm"], r["cpm"], 0.006)
            _close((g["labor_spend"] or 0) + (g["parts_spend"] or 0),
                   r["spend"])
        _close(res["totals"]["spend"], rep["total_spend"])
        assert "cpm" not in res["totals"]           # CPM no se suma
        assert res["rows"][0]["period"] != ""
    # MEM 77 junio: 4,000 mi; gasto junio = 380 + 257 = 637.
    rep = reports.cpm_report("2026-06-01", "2026-06-30")
    mem = next(r for r in rep["by_unit"] if r["unit"] == "MEM 77")
    assert mem["miles"] == 4000 and mem["spend"] == 637.0
    print("OK paridad: cpm == cpm_report (+ split labor/parts por unidad)")


def test_export_csv_xlsx():
    _seed()
    workorders.create_wo("=HYPERLINK(\"x\")", "+cmd injection")
    res = _run({"source": "work_orders",
                "columns": [{"field": "unit"}, {"field": "title"},
                            {"field": "total"}, {"field": "status"},
                            {"field": "created_at"},
                            {"field": "waiting_parts"}],
                "totals": True})
    data, mime, fname = report_export.render(res, "csv", "Open WOs")
    assert mime.startswith("text/csv") and fname.endswith(".csv")
    text = data.decode("utf-8")
    assert text.startswith("﻿")
    assert "'=HYPERLINK" in text and "'+cmd injection" in text
    assert "In progress" in text and "Total" in text.splitlines()[-1]
    data, mime, fname = report_export.render(res, "xlsx", "Open WOs",
                                             "All time")
    assert fname.endswith(".xlsx") and "spreadsheetml" in mime
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "Open WOs" and ws["A4"].value == "Unit"
    vals = [c.value for c in ws["A"]]
    assert '=HYPERLINK("x")' in vals
    cell = next(c for c in ws["A"] if c.value == '=HYPERLINK("x")')
    assert cell.data_type == "s"
    _raises(report_export.render, "format must be csv or xlsx", res, "pdf",
            "x")
    print("OK export CSV (BOM, sin fórmulas) y XLSX (strings, formatos)")


if __name__ == "__main__":
    test_registry_and_catalog()
    test_validation_errors()
    test_presets()
    test_filters_aggregation_sort_totals()
    test_text_pushdown_autoescape()
    test_time_filters_and_derived()
    test_pii_masking_and_tenant_isolation()
    test_library_validates_and_runs()
    test_spend_parity()
    test_cpm_parity()
    test_export_csv_xlsx()
    print("\nALL REPORT ENGINE TESTS PASSED")
